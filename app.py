from __future__ import annotations

import csv
import io
import os
import random
import secrets
import shutil
import sqlite3
try:
    import psycopg
    from psycopg.rows import dict_row
except ImportError:
    psycopg = None
from datetime import date, datetime
from zoneinfo import ZoneInfo
from functools import wraps
from pathlib import Path

import requests
from dotenv import load_dotenv
from flask import Flask, flash, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook, load_workbook
from werkzeug.security import check_password_hash, generate_password_hash

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / '.env')
DATA_DIR = BASE_DIR / 'data'
BACKUP_DIR = BASE_DIR / 'backups'
DATA_DIR.mkdir(exist_ok=True)
BACKUP_DIR.mkdir(exist_ok=True)
DB_PATH = Path(os.getenv('DB_PATH', DATA_DIR / 'auguri_v3.db'))

app = Flask(__name__)
app.secret_key = os.getenv('SECRET_KEY', secrets.token_hex(32))
app.config['MAX_CONTENT_LENGTH'] = 5 * 1024 * 1024


def using_postgres():
    return bool(os.getenv('DATABASE_URL', '').strip())


class PgCursorCompat:
    def __init__(self, cur): self.cur = cur
    def fetchone(self): return self.cur.fetchone()
    def fetchall(self): return self.cur.fetchall()


class PgCompat:
    def __init__(self, con): self.con = con
    def execute(self, sql, params=()):
        sql = sql.replace('?', '%s')
        cur = self.con.cursor(row_factory=dict_row)
        cur.execute(sql, params)
        return PgCursorCompat(cur)
    def executescript(self, script):
        cur = self.con.cursor()
        cur.execute(script)
    def commit(self): self.con.commit()
    def close(self): self.con.close()


def get_db():
    url = os.getenv('DATABASE_URL', '').strip()
    if url:
        if psycopg is None:
            raise RuntimeError('psycopg non installato')
        return PgCompat(psycopg.connect(url, autocommit=False))
    con = sqlite3.connect(DB_PATH, timeout=30)
    con.row_factory = sqlite3.Row
    con.execute('PRAGMA foreign_keys=ON')
    return con

def init_db():
    con = get_db()
    if using_postgres():
        con.executescript("""
        CREATE TABLE IF NOT EXISTS contacts (
          id BIGSERIAL PRIMARY KEY, name TEXT NOT NULL, phone TEXT NOT NULL, birth_date TEXT NOT NULL,
          category TEXT NOT NULL DEFAULT 'Generale', notes TEXT NOT NULL DEFAULT '', enabled INTEGER NOT NULL DEFAULT 1,
          created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS templates (
          id BIGSERIAL PRIMARY KEY, label TEXT NOT NULL, category TEXT NOT NULL DEFAULT 'Generale', template_name TEXT NOT NULL,
          language TEXT NOT NULL DEFAULT 'it', preview_text TEXT NOT NULL, enabled INTEGER NOT NULL DEFAULT 1,
          created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS sent_log (
          id BIGSERIAL PRIMARY KEY, contact_id BIGINT REFERENCES contacts(id) ON DELETE SET NULL, template_id BIGINT REFERENCES templates(id) ON DELETE SET NULL,
          recipient_name TEXT NOT NULL, phone TEXT NOT NULL, template_name TEXT NOT NULL, rendered_text TEXT NOT NULL, sent_at TEXT NOT NULL,
          status TEXT NOT NULL, provider_message_id TEXT DEFAULT '', response_excerpt TEXT DEFAULT '');
        CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT NOT NULL);
        """)
    else:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS contacts (id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,phone TEXT NOT NULL,birth_date TEXT NOT NULL,category TEXT NOT NULL DEFAULT 'Generale',notes TEXT NOT NULL DEFAULT '',enabled INTEGER NOT NULL DEFAULT 1,created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS templates (id INTEGER PRIMARY KEY AUTOINCREMENT,label TEXT NOT NULL,category TEXT NOT NULL DEFAULT 'Generale',template_name TEXT NOT NULL,language TEXT NOT NULL DEFAULT 'it',preview_text TEXT NOT NULL,enabled INTEGER NOT NULL DEFAULT 1,created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS sent_log (id INTEGER PRIMARY KEY AUTOINCREMENT,contact_id INTEGER,template_id INTEGER,recipient_name TEXT NOT NULL,phone TEXT NOT NULL,template_name TEXT NOT NULL,rendered_text TEXT NOT NULL,sent_at TEXT NOT NULL,status TEXT NOT NULL,provider_message_id TEXT DEFAULT '',response_excerpt TEXT DEFAULT '',FOREIGN KEY(contact_id) REFERENCES contacts(id) ON DELETE SET NULL,FOREIGN KEY(template_id) REFERENCES templates(id) ON DELETE SET NULL);
        CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY,value TEXT NOT NULL);
        """)
    # Additive migration: existing Neon and SQLite contacts remain untouched.
    if using_postgres():
        con.execute('ALTER TABLE contacts ADD COLUMN IF NOT EXISTS auto_birthday INTEGER NOT NULL DEFAULT 0')
    else:
        cols = [r['name'] for r in con.execute('PRAGMA table_info(contacts)').fetchall()]
        if 'auto_birthday' not in cols:
            con.execute('ALTER TABLE contacts ADD COLUMN auto_birthday INTEGER NOT NULL DEFAULT 0')
    defaults={'send_hour':'09','send_minute':'00','auto_send':'1','default_country':'+39','admin_password_hash':'','app_title':'Auguri WhatsApp'}
    for k,v in defaults.items():
        con.execute('INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO NOTHING',(k,v))
    con.commit(); con.close()

def setting(key: str, default='') -> str:
    con = get_db()
    row = con.execute('SELECT value FROM settings WHERE key=?', (key,)).fetchone()
    con.close()
    return row['value'] if row else default


def set_setting(key: str, value: str):
    con = get_db()
    con.execute('INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value', (key, str(value)))
    con.commit(); con.close()


def normalize_phone(phone: str) -> str:
    raw = ''.join(ch for ch in str(phone).strip() if ch.isdigit() or ch == '+')
    if raw.startswith('+'):
        return raw
    prefix = setting('default_country', '+39')
    return prefix + raw.lstrip('0')


def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not session.get('logged_in'):
            return redirect(url_for('login', next=request.path))
        return fn(*args, **kwargs)
    return wrapped


def admin_password_hash() -> str:
    db_hash = setting('admin_password_hash', '')
    if db_hash:
        return db_hash
    env_password = os.getenv('ADMIN_PASSWORD', '')
    if env_password:
        hashed = generate_password_hash(env_password)
        set_setting('admin_password_hash', hashed)
        return hashed
    return ''


def birthday_today(birth_date: str) -> bool:
    try:
        dt = datetime.strptime(birth_date, '%Y-%m-%d')
        now = datetime.now()
        return (dt.month, dt.day) == (now.month, now.day)
    except Exception:
        return False


def age_today(birth_date: str):
    try:
        born = datetime.strptime(birth_date, '%Y-%m-%d').date()
        today = date.today()
        return today.year - born.year - ((today.month, today.day) < (born.month, born.day))
    except Exception:
        return ''


def due_contacts(automatic_only=False):
    con = get_db()
    rows = con.execute('SELECT * FROM contacts WHERE enabled=1' + (' AND auto_birthday=1' if automatic_only else '') + ' ORDER BY name').fetchall()
    con.close()
    return [r for r in rows if birthday_today(r['birth_date'])]


def already_sent_today(contact_id: int) -> bool:
    today = datetime.now().strftime('%Y-%m-%d')
    con = get_db()
    row = con.execute("SELECT 1 FROM sent_log WHERE contact_id=? AND substr(sent_at,1,10)=? AND status='OK' LIMIT 1", (contact_id, today)).fetchone()
    con.close()
    return bool(row)


def choose_template(category: str):
    con = get_db()
    rows = con.execute('SELECT * FROM templates WHERE enabled=1 AND category=? ORDER BY id', (category,)).fetchall()
    if not rows:
        rows = con.execute("SELECT * FROM templates WHERE enabled=1 AND category='Generale' ORDER BY id").fetchall()
    if not rows:
        rows = con.execute('SELECT * FROM templates WHERE enabled=1 ORDER BY id').fetchall()
    con.close()
    return random.choice(rows) if rows else None


def render_preview(text: str, contact) -> str:
    age = age_today(contact['birth_date'])
    return (text or '').replace('{nome}', contact['name']).replace('{eta}', str(age))


def whatsapp_config():
    return {
      'token': os.getenv('WHATSAPP_TOKEN', '').strip(),
      'phone_number_id': os.getenv('WHATSAPP_PHONE_NUMBER_ID', '').strip(),
      'api_version': os.getenv('WHATSAPP_API_VERSION', 'v26.0').strip(),
    }


def send_whatsapp_template(contact, template_row):
    cfg = whatsapp_config()
    if not cfg['token'] or not cfg['phone_number_id']:
        raise RuntimeError('WhatsApp Cloud API non configurata: impostare WHATSAPP_TOKEN e WHATSAPP_PHONE_NUMBER_ID nel file .env')

    # Questa V3 usa due parametri BODY standard: {{1}}=nome e {{2}}=età.
    # Se il template Meta usa una struttura diversa, modificare components in base al template approvato.
    payload = {
      'messaging_product': 'whatsapp',
      'to': normalize_phone(contact['phone']).replace('+', ''),
      'type': 'template',
      'template': {
        'name': template_row['template_name'],
        'language': {'code': template_row['language']},
        'components': [
          {
            'type': 'body',
            'parameters': [
              {'type': 'text', 'text': contact['name']},
              {'type': 'text', 'text': str(age_today(contact['birth_date']))},
            ]
          }
        ]
      }
    }
    url = f"https://graph.facebook.com/{cfg['api_version']}/{cfg['phone_number_id']}/messages"
    response = requests.post(url, json=payload, headers={'Authorization': f"Bearer {cfg['token']}", 'Content-Type': 'application/json'}, timeout=30)
    excerpt = response.text[:1000]
    if not response.ok:
        raise RuntimeError(f'HTTP {response.status_code}: {excerpt}')
    data = response.json()
    msg_id = ''
    try:
        msg_id = data.get('messages', [{}])[0].get('id', '')
    except Exception:
        pass
    return msg_id, excerpt


def log_send(contact, template_row, rendered_text: str, status: str, provider_message_id='', response_excerpt=''):
    con = get_db()
    con.execute('''INSERT INTO sent_log(contact_id,template_id,recipient_name,phone,template_name,rendered_text,sent_at,status,provider_message_id,response_excerpt)
                   VALUES(?,?,?,?,?,?,?,?,?,?)''',
                (contact['id'], template_row['id'] if template_row else None, contact['name'], contact['phone'], template_row['template_name'] if template_row else '', rendered_text,
                 datetime.now().strftime('%Y-%m-%d %H:%M:%S'), status, provider_message_id, response_excerpt))
    con.commit(); con.close()


def process_birthdays(force=False, automatic_only=False):
    results = []
    for contact in due_contacts(automatic_only=automatic_only):
        if already_sent_today(contact['id']) and not force:
            results.append((contact['name'], 'GIÀ INVIATO'))
            continue
        template = choose_template(contact['category'])
        if not template:
            results.append((contact['name'], 'NESSUN TEMPLATE'))
            continue
        rendered = render_preview(template['preview_text'], contact)
        try:
            msg_id, raw = send_whatsapp_template(contact, template)
            log_send(contact, template, rendered, 'OK', msg_id, raw)
            results.append((contact['name'], 'OK'))
        except Exception as exc:
            log_send(contact, template, rendered, 'ERRORE', '', str(exc)[:1000])
            results.append((contact['name'], f'ERRORE: {exc}'))
    return results


@app.get('/login')
def login():
    return render_template('login.html', configured=bool(admin_password_hash()))


@app.post('/login')
def login_post():
    stored = admin_password_hash()
    if not stored:
        flash('Password amministratore non configurata. Impostare ADMIN_PASSWORD nel file .env e riavviare.', 'error')
        return redirect(url_for('login'))
    if check_password_hash(stored, request.form.get('password', '')):
        session['logged_in'] = True
        return redirect(request.args.get('next') or url_for('index'))
    flash('Password non corretta.', 'error')
    return redirect(url_for('login'))


@app.get('/logout')
def logout():
    session.clear(); return redirect(url_for('login'))


@app.get('/')
@login_required
def index():
    con = get_db()
    contacts = con.execute('SELECT * FROM contacts ORDER BY substr(birth_date,6,5), name').fetchall()
    templates = con.execute('SELECT * FROM templates ORDER BY category,label').fetchall()
    logs = con.execute('SELECT * FROM sent_log ORDER BY id DESC LIMIT 100').fetchall()
    categories = [(r['category'] if hasattr(r, 'keys') else r[0]) for r in con.execute("SELECT category FROM contacts UNION SELECT category FROM templates ORDER BY 1").fetchall() if (r['category'] if hasattr(r, 'keys') else r[0])]
    con.close()
    due = due_contacts()
    cfg = whatsapp_config()
    stats = {
      'today': len(due),
      'contacts': len(contacts),
      'templates': len(templates),
      'sent_month': sum(1 for l in logs if l['status']=='OK' and l['sent_at'].startswith(datetime.now().strftime('%Y-%m'))),
    }
    sets = {k: setting(k) for k in ['send_hour','send_minute','auto_send','default_country','app_title']}
    upcoming = sorted(contacts, key=lambda c: ((date(date.today().year, int(c['birth_date'][5:7]), min(int(c['birth_date'][8:10]), 28)) - date.today()).days % 365) if len(c['birth_date']) >= 10 else 999)[:10]
    return render_template('index.html', upcoming=upcoming, contacts=contacts, templates=templates, logs=logs, categories=categories, due=due, age_today=age_today, stats=stats, settings=sets, whatsapp_ready=bool(cfg['token'] and cfg['phone_number_id']), api_version=cfg['api_version'])


@app.post('/contacts/add')
@login_required
def contacts_add():
    f = request.form
    if not f.get('name') or not f.get('phone') or not f.get('birth_date'):
        flash('Nome, telefono e data di nascita sono obbligatori.', 'error'); return redirect(url_for('index'))
    con=get_db(); con.execute('INSERT INTO contacts(name,phone,birth_date,category,notes,auto_birthday) VALUES(?,?,?,?,?,?)',
      (f['name'].strip(), normalize_phone(f['phone']), f['birth_date'], f.get('category','Generale').strip() or 'Generale', f.get('notes','').strip(),1 if f.get('auto_birthday') else 0))
    con.commit(); con.close(); flash('Contatto aggiunto.', 'ok'); return redirect(url_for('index'))


@app.post('/contacts/<int:cid>/toggle')
@login_required
def contacts_toggle(cid):
    con=get_db(); con.execute('UPDATE contacts SET enabled=1-enabled WHERE id=?',(cid,)); con.commit(); con.close(); return redirect(url_for('index'))


@app.post('/contacts/<int:cid>/auto-toggle')
@login_required
def contacts_auto_toggle(cid):
    con=get_db()
    con.execute('UPDATE contacts SET auto_birthday=1-auto_birthday WHERE id=?',(cid,))
    con.commit(); con.close()
    flash('Preferenza invio automatico aggiornata.', 'ok')
    return redirect(url_for('index'))


@app.post('/contacts/<int:cid>/delete')
@login_required
def contacts_delete(cid):
    con=get_db(); con.execute('DELETE FROM contacts WHERE id=?',(cid,)); con.commit(); con.close(); flash('Contatto eliminato.','ok'); return redirect(url_for('index'))


@app.post('/contacts/<int:cid>/edit')
@login_required
def contacts_edit(cid):
    f=request.form; con=get_db(); con.execute('UPDATE contacts SET name=?,phone=?,birth_date=?,category=?,notes=? WHERE id=?',
      (f['name'].strip(), normalize_phone(f['phone']), f['birth_date'], f.get('category','Generale').strip() or 'Generale', f.get('notes','').strip(), cid))
    con.execute('UPDATE contacts SET auto_birthday=? WHERE id=?',(1 if f.get('auto_birthday') else 0,cid))
    con.commit(); con.close(); flash('Contatto aggiornato.','ok'); return redirect(url_for('index'))


@app.post('/templates/add')
@login_required
def templates_add():
    f=request.form
    fields=[f.get('label','').strip(), f.get('template_name','').strip(), f.get('preview_text','').strip()]
    if not all(fields):
        flash('Etichetta, nome template Meta e anteprima sono obbligatori.','error'); return redirect(url_for('index'))
    con=get_db(); con.execute('INSERT INTO templates(label,category,template_name,language,preview_text) VALUES(?,?,?,?,?)',
      (fields[0], f.get('category','Generale').strip() or 'Generale', fields[1], f.get('language','it').strip() or 'it', fields[2]))
    con.commit(); con.close(); flash('Template aggiunto.','ok'); return redirect(url_for('index'))


@app.post('/templates/<int:tid>/toggle')
@login_required
def templates_toggle(tid):
    con=get_db(); con.execute('UPDATE templates SET enabled=1-enabled WHERE id=?',(tid,)); con.commit(); con.close(); return redirect(url_for('index'))


@app.post('/templates/<int:tid>/delete')
@login_required
def templates_delete(tid):
    con=get_db(); con.execute('DELETE FROM templates WHERE id=?',(tid,)); con.commit(); con.close(); flash('Template eliminato.','ok'); return redirect(url_for('index'))


@app.post('/templates/<int:tid>/edit')
@login_required
def templates_edit(tid):
    f=request.form; con=get_db(); con.execute('UPDATE templates SET label=?,category=?,template_name=?,language=?,preview_text=? WHERE id=?',
      (f['label'].strip(), f.get('category','Generale').strip() or 'Generale', f['template_name'].strip(), f.get('language','it').strip() or 'it', f['preview_text'].strip(), tid))
    con.commit(); con.close(); flash('Template aggiornato.','ok'); return redirect(url_for('index'))


@app.get('/manual/<int:cid>')
@login_required
def manual_whatsapp(cid):
    con = get_db()
    contact = con.execute('SELECT * FROM contacts WHERE id=?', (cid,)).fetchone()
    con.close()
    if not contact:
        flash('Contatto non trovato.', 'error')
        return redirect(url_for('index'))
    template = choose_template(contact['category'])
    if not template:
        flash('Nessuna frase attiva disponibile.', 'error')
        return redirect(url_for('index'))
    return render_template('manual_whatsapp.html', contact=contact, template=template,
                           message=render_preview(template['preview_text'], contact),
                           phone=normalize_phone(contact['phone']))


@app.post('/send/<int:cid>')
@login_required
def send_one(cid):
    con=get_db(); c=con.execute('SELECT * FROM contacts WHERE id=?',(cid,)).fetchone(); con.close()
    if not c: flash('Contatto non trovato.','error'); return redirect(url_for('index'))
    t=choose_template(c['category'])
    if not t: flash('Nessun template attivo disponibile.','error'); return redirect(url_for('index'))
    rendered=render_preview(t['preview_text'], c)
    try:
        msg_id, raw=send_whatsapp_template(c,t); log_send(c,t,rendered,'OK',msg_id,raw); flash(f'Auguri inviati a {c["name"]}.','ok')
    except Exception as exc:
        log_send(c,t,rendered,'ERRORE','',str(exc)); flash(f'Invio non riuscito: {exc}','error')
    return redirect(url_for('index'))


@app.post('/send-today')
@login_required
def send_today():
    results=process_birthdays(force=request.form.get('force')=='1')
    if not results: flash('Nessun compleanno oggi.','info')
    else: flash(' | '.join(f'{n}: {s}' for n,s in results),'ok')
    return redirect(url_for('index'))


@app.post('/settings')
@login_required
def settings_save():
    for key in ['send_hour','send_minute','default_country','app_title']:
        set_setting(key, request.form.get(key, setting(key)))
    set_setting('auto_send','1' if request.form.get('auto_send') else '0')
    flash('Impostazioni salvate.','ok'); return redirect(url_for('index'))


@app.post('/password')
@login_required
def password_change():
    p=request.form.get('password','')
    if len(p)<8: flash('La password deve avere almeno 8 caratteri.','error')
    else: set_setting('admin_password_hash',generate_password_hash(p)); flash('Password aggiornata.','ok')
    return redirect(url_for('index'))


@app.post('/import')
@login_required
def import_contacts():
    up=request.files.get('file')
    if not up or not up.filename: flash('Seleziona un file CSV o XLSX.','error'); return redirect(url_for('index'))
    ext=Path(up.filename).suffix.lower(); rows=[]
    try:
        if ext=='.csv':
            text=io.StringIO(up.stream.read().decode('utf-8-sig')); rows=list(csv.DictReader(text))
        elif ext=='.xlsx':
            wb=load_workbook(up.stream,data_only=True); ws=wb.active; headers=[str(c.value or '').strip().lower() for c in ws[1]]
            for vals in ws.iter_rows(min_row=2,values_only=True): rows.append({headers[i]: vals[i] for i in range(min(len(headers),len(vals)))})
        else: raise ValueError('Formato non supportato')
        con=get_db(); count=0
        for r in rows:
            name=str(r.get('nome') or r.get('name') or '').strip(); phone=str(r.get('telefono') or r.get('phone') or '').strip(); bd=r.get('data_nascita') or r.get('birth_date') or ''
            if isinstance(bd, datetime): bd=bd.strftime('%Y-%m-%d')
            bd=str(bd).strip()
            cat=str(r.get('categoria') or r.get('category') or 'Generale').strip() or 'Generale'; notes=str(r.get('note') or r.get('notes') or '').strip()
            if name and phone and bd:
                con.execute('INSERT INTO contacts(name,phone,birth_date,category,notes,auto_birthday) VALUES(?,?,?,?,?,?)',(name,normalize_phone(phone),bd,cat,notes, 1 if str(r.get('invio_automatico') or '').strip().lower() in ('1','si','sì','true','yes') else 0)); count+=1
        con.commit(); con.close(); flash(f'Importati {count} contatti.','ok')
    except Exception as exc: flash(f'Importazione fallita: {exc}','error')
    return redirect(url_for('index'))


@app.get('/export.xlsx')
@login_required
def export_contacts():
    con=get_db(); rows=con.execute('SELECT name,phone,birth_date,category,notes,enabled,auto_birthday FROM contacts ORDER BY name').fetchall(); con.close()
    wb=Workbook(); ws=wb.active; ws.title='Contatti'; ws.append(['nome','telefono','data_nascita','categoria','note','attivo'])
    for r in rows: ws.append([r[k] for k in ['name','phone','birth_date','category','notes','enabled']])
    bio=io.BytesIO(); wb.save(bio); bio.seek(0)
    return send_file(bio,as_attachment=True,download_name='contatti_auguri.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.get('/modello.xlsx')
@login_required
def model_xlsx():
    wb=Workbook(); ws=wb.active; ws.title='Contatti'; ws.append(['nome','telefono','data_nascita','categoria','note']); ws.append(['Mario Rossi','3331234567','1980-09-10','Amici','Esempio'])
    bio=io.BytesIO(); wb.save(bio); bio.seek(0)
    return send_file(bio,as_attachment=True,download_name='modello_importazione_auguri.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.get('/backup')
@login_required
def backup_db():
    if using_postgres():
        flash('Con Neon/PostgreSQL il backup SQLite locale non è utilizzato. Usa Export Excel per la rubrica.', 'info')
        return redirect(url_for('index'))
    target=BACKUP_DIR / f"auguri_v3_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"; shutil.copy2(DB_PATH,target)
    return send_file(target,as_attachment=True,download_name=target.name)


@app.post('/simulate')
@login_required
def simulate():
    """Read-only dry run; never calls Meta and never writes to the send log."""
    rows=[]
    for contact in due_contacts():
        template=choose_template(contact['category'])
        rows.append({'contact':contact['name'], 'category':contact['category'],
                     'template':template['template_name'] if template else None,
                     'preview':render_preview(template['preview_text'],contact) if template else None,
                     'already_sent':already_sent_today(contact['id']), 'automatic_enabled':bool(contact['auto_birthday'])})
    return {'simulation':True,'whatsapp_called':False,'records_written':False,
            'date_italy':datetime.now(ZoneInfo('Europe/Rome')).date().isoformat(),
            'results':rows}


@app.get('/health')
def health():
    return {'ok': True, 'time': datetime.now().isoformat(timespec='seconds')}


@app.post('/cron/run')
def cron_http():
    secret=os.getenv('CRON_SECRET','')
    if not secret or request.headers.get('X-Cron-Secret') != secret:
        return {'ok': False, 'error': 'unauthorized'}, 401
    if setting('auto_send','1')!='1': return {'ok':True,'skipped':'auto_send disabled'}
    now=datetime.now(ZoneInfo('Europe/Rome')); hh=int(setting('send_hour','09')); mm=int(setting('send_minute','00'))
    # Permette una finestra di 15 minuti: ideale se il cron gira ogni 5 minuti.
    current=now.hour*60+now.minute; target=hh*60+mm
    if not (target <= current <= target+14): return {'ok':True,'skipped':'outside_send_window','italy_time':now.strftime('%H:%M'),'scheduled_time':f'{hh:02d}:{mm:02d}'}
    return {'ok':True,'results':process_birthdays(force=False, automatic_only=True)}


init_db()

if __name__ == '__main__':
    app.run(host='127.0.0.1', port=int(os.getenv('PORT','5000')), debug=os.getenv('FLASK_DEBUG','0')=='1')
